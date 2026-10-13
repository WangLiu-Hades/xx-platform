from __future__ import annotations
import json, re, hashlib
from pathlib import Path
from datetime import datetime, timezone
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[3]
OUT=ROOT/'docs/project-governance'; AUDIT=OUT/'audit'
SOURCE=ROOT/'docs/requirements/source/多维数据分析平台开发模块与人员分工.docx'
OUT.mkdir(parents=True,exist_ok=True)
api=json.loads((AUDIT/'api-inventory.json').read_text('utf-8'))
inventory=json.loads((AUDIT/'workspace-inventory-summary.json').read_text('utf-8'))

def write(name,text): (OUT/name).write_text(text.strip()+'\n',encoding='utf-8')

# ID, module delivery coordinator, name, API implementation, data implementation, frontend, acceptance
MODULES=[
('M01','DEV-01','平台底座','统一响应、会话/RBAC、审计、文件、任务、字典、Context、幂等和迁移发布','评审数据输入快照与Worker领取/发布协议','API客户端、启动配置与无权/过期处理','T01/T02/T05；旧会话无效、越权及重放拒绝'),
('M02','DEV-02','前端应用壳','提供会话、菜单、权限与client-config','提供脱敏样例','维护现有Vue工程；所有业务页面、组件和HTTP适配均归DEV-02','28菜单可访问；Mock/HTTP分离；刷新/失败可恢复'),
('M03','DEV-03','接入标准化','批次/回放管理API、权限、任务提交','源适配器、映射/时区/坐标、校验隔离、稳定去重、标准事实写入、质量汇总','按已确认范围展示接入质量；当前无专门接入管理页','T19；同批次回放不重计；错误可定位到源记录'),
('M04','DEV-02','系统首页','系统首页聚合响应、权限与缓存边界','共用指标定义和快照计算','仅系统首页12类区块、钻取、局部失败','T03；采集量与入库量分开；指标asOf/单位一致'),
('M05','DEV-03','多源检索','公开检索/详情/搜图API、范围约束、分页、导出编排','SQL查询投影、索引计划、四源字段语义、人像模型适配与效果验证','四类搜索、详情、分页、受控搜图入口','T04/T06；类型/场所隔离；人像阈值须签认'),
('M06','DEV-01','通用研判动作','Context、权限重校、解密申请、导出与任务编排','统计、轨迹等可复用查询/计算实现','右键动作和跨页上下文；复制/重置是浏览器动作','T05/T08/T09；过期/缩权拒绝；不传展示号当内部ID'),
('M07','DEV-02','地图研判','地图/轨迹/热力/区域摘要API与限域限量','WGS84投影、空间查询、聚合、按目标排序轨迹','OpenLayers图层、框选/单选、列表同步、轨迹/热力与重置','T08；多点选择正确；不同目标不连线'),
('M08','DEV-03','研判任务','领域任务创建/修订、权限、统一任务执行与结果发布','伴随/碰撞/区域计算、证据、参数和算法版本','三类表单、进度/结果与二次动作','T09；取消/失败不发布半成品；覆盖率不解释为嫌疑概率'),
('M09','DEV-03','数据集运算','数据集资源/成员/版本/范围、引用保护、任务API','导入解析、固定版本快照、交并差及表达式执行','导入、选择/表达式预览、结果数据集','T10；CAPTURE按captureId、TARGET按类型+ID；禁止异型运算'),
('M10','DEV-01','布控与通知','规则生命周期、告警事务/去重账本、计划调度、通道适配/投递/回执','按签认规则计算命中候选及证据；不直接发短信','综合布控、告警、运维通知、平台消息入口','T11/T17；撤销边界、UNKNOWN回执对账、频控不重发'),
('M11','DEV-03','人员库档案模型','人员库CRUD/授权、档案修订/合并拆分事务、谱系、版本、模型目录API','导入解析、身份候选/置信度/证据、常驻模型执行及验证','库/人员、图码/车码档案、数据态势；未绑定高级接口单独排期','T07/T12；P-016/P-017关闭后验真实规则；不覆盖原事实'),
('M12','DEV-01','运维后台','场所/设备/分组、设备命令协议适配、逐台回执、报表任务及下载','设备遥测标准化、历史采集聚合、报表数据计算/文件生成','管理、参数配置、报表表单/详情','T14/T16；成功设备免重发；历史报表冻结；Excel/ZIP正确'),
('M13','DEV-02','运维可视化','运维指标/设备地图/报表预览公开API','复用M04指标定义及M12遥测/报表投影','运维首页四类区块、设备分布、状态/热力和联动','T13/T15；不再实现第二套运维首页/指标公式'),
('M14','DEV-01','设置与审计','用户/角色、权限授予、禁用/删除、日志查询/导出与脱敏','提供受限查询需要的范围谓词，不维护另一套RBAC','用户/角色/日志界面和授权呈现','T02/T18；普通用户只看自身；删除引用保护'),
('M15','DEV-04','质量交付','DEV-01修复后端缺陷并提供运行探针','DEV-03修复数据缺陷并提供可重复样本','DEV-02修复前端缺陷并维护定位标识','T01–T27、CI/CD、真实集成/容量/恢复；业务方签最终业务验收')]

WBS=[
('W01','基线与运行骨架','M01/M15','无',[(6,8),(3,5),(2,3),(3,5)]),
('W02','会话与权限闭环','M01/M02/M14','W01；P-001/P-002/P-021',[(10,14),(1,2),(4,6),(5,7)]),
('W03','任务文件Context底座','M01/M06','W01；P-008/P-009',[(12,18),(4,6),(4,6),(6,8)]),
('W04','四源接入与回放','M03','W01/W03协议；P-004/P-005',[(5,8),(15,22),(0,1),(5,8)]),
('W05','检索与GIS真实链路','M05/M07','W02/W03/W04；P-007/P-010',[(8,12),(12,18),(5,8),(7,10)]),
('W06','管理资源和版本','M09/M11/M12/M14','W02/W03；P-013/P-016已确定部分',[(10,14),(6,10),(4,6),(6,9)]),
('W07','指标首页和八类报表','M04/M12/M13','W04/W06；P-006/P-020',[(8,12),(10,15),(3,5),(6,9)]),
('W08','研判和集合运算','M08/M09','W03/W05/W06；P-011/P-013',[(8,12),(15,22),(4,6),(7,10)]),
('W09','人员库身份档案常驻模型','M11','W03/W04/W06；P-016/P-017',[(10,15),(20,30),(5,8),(8,12)]),
('W10','布控命中和可靠通知','M10','W02/W03/W04；P-014/P-015',[(10,16),(8,12),(3,5),(7,10)]),
('W11','真实设备命令与回执','M12','W03/W06；P-018/P-019',[(10,16),(4,6),(3,5),(7,10)]),
('W12','生产加固验收移交','M15及全模块','前述已纳入范围的工作包；P-003/Q-002',[(8,12),(6,10),(4,6),(12,18)])]
totals=[tuple(sum(w[4][i][j] for w in WBS) for j in range(2)) for i in range(4)]
(AUDIT/'work-packages.json').write_text(json.dumps({'roles':['DEV-01','DEV-03','DEV-02','DEV-04'],'packages':[{'id':w[0],'name':w[1],'modules':w[2],'dependencies':w[3],'personDays':dict(zip(['DEV-01','DEV-03','DEV-02','DEV-04'],w[4]))} for w in WBS],'totals':dict(zip(['DEV-01','DEV-03','DEV-02','DEV-04'],totals))},ensure_ascii=False,indent=2),encoding='utf-8')

pending_rows=[]
for line in (ROOT/'contracts/api/OPENAPI_PENDING.md').read_text('utf-8').splitlines():
 if re.match(r'\| P-\d{3} \|',line):pending_rows.append([s.strip() for s in line.split('|')[1:-1]])
owners=[('DEV-01','业务/安全负责人','认证方案ADR、会话生命周期和部署边界'),('DEV-01','产品/安全负责人','角色×操作×数据/场所范围矩阵及缩权用例'),('DEV-04汇总、DEV-01/03估算','业务/运维负责人','数据分布、日增、峰值、保留期、测量环境和SLO'),('DEV-03','设备厂商/数据负责人','四类采集+设备状态原始样本、协议/字段/时间/坐标/去重映射'),('DEV-03；DEV-01审查权限','数据/安全负责人','回放/补数范围、幂等与补偿用例'),('DEV-03','业务负责人','逐项指标公式、时间边界、去重、单位、源采集与入库口径'),('DEV-03或算法供应方','业务/算法/安全负责人','授权模型、标注评测集、阈值/候选上限、用途及导出边界'),('DEV-01','安全负责人','密钥/证书、审批、明文传递/保留/下载规则'),('DEV-01；DEV-03测量','业务/数据/安全负责人','查询取消、限流、Context TTL和导出限值'),('DEV-02；DEV-03坐标','产品/采购/数据负责人','坐标转换依据、底图许可、限额与菜单确认'),('DEV-03或算法供应方','业务/算法负责人','时空规则、可解释证据、版本及指标含义'),('DEV-03','已收敛限定范围','ID相等规则和表达式语法；不包括跨源身份拟合'),('DEV-01；DEV-03成员模型','业务/安全负责人','共享、继承权限、引用删除、模板/版本生命周期'),('DEV-03匹配；DEV-01告警','业务/算法负责人','命中和去重规则、运行中流水/结束后统计及撤销边界'),('DEV-01','运营/通道供应方','通道协议/账号、模板、频控、回执/UNKNOWN对账与人工接管'),('DEV-03；DEV-01事务/权限','业务/算法/安全负责人','身份等价、冲突/合并拆分、证据、谱系、人工权限及评测集'),('DEV-03或算法供应方','业务/算法负责人','仅RESIDENCE的输入、参数、阈值、标签、版本和评测标准'),('DEV-01；DEV-03遥测','设备厂商/运维负责人','参数/命令目录、协议、回执、批量重试和实机'),('DEV-01','安全/运维负责人','高风险审批、二次确认与生效规则'),('DEV-03公式；DEV-01任务','业务/运维负责人','八类报表公式、异常识别、保留和失败重试'),('DEV-01','安全/业务负责人','用户生命周期、密码/日志留存脱敏及纯UI动作审计范围')]
assert len(pending_rows)==len(owners)==21
pending_table='\n'.join(f'| {r[0]} | {"已收敛，仅ID集合语法" if r[0]=="P-012" else "仍有未决规则"} | {o[0]} | {o[1]} | {o[2]} |' for r,o in zip(pending_rows,owners))
write('PENDING_REGISTER.md',f'''
# 待确认事项定位与关闭责任

核对日期：2026-10-10。P-016原句位于 `../../contracts/database/data-dictionary.md` 第1895行；相同约束也写入 `schema.sql` 对 `analysis.profile_revision` 的注释。它是设计阻塞标记，不是数据库错误或已经完成的算法。

| 文件 | 地位和用途 |
| --- | --- |
| [OPENAPI_PENDING.md](../../contracts/api/OPENAPI_PENDING.md) | P-001–P-021的业务决定主清单；接口/字段pending应引用这里 |
| [pending-design.md](../../contracts/database/pending-design.md) | 数据库视角投影；不能独立关闭同编号问题 |
| [11_架构风险与待确认事项.md](../system-design/11_架构风险与待确认事项.md) | ISSUE-001–016源资料冲突、Q-001–008缺失输入；编号与P系列不是同一清单 |
| [openapi.yaml](../../contracts/api/openapi.yaml) | 机器契约中的x-design-status和x-pending-reason，决定哪些操作/字段仍未定 |
| [data-dictionary.md](../../contracts/database/data-dictionary.md) | 把未决规则标到表/字段；不作为业务签认主账本 |
| [api-storage-map.md](../../contracts/database/api-storage-map.md) | 需求→API→存储的影响追踪，供改规则时找到受影响资源 |

目前21个编号中P-012的ID集合语法已收敛，20项仍含未决决定。65是**操作级pending数量**，不是待确认事项数量；字段和项目级阻塞需一起检查。下表的负责人是本次修订后的技术提案责任，决策人员是所需角色，不代表已任命、已签认或已提供资料。

| ID | 状态 | 技术提案责任 | 需确认角色 | 最小关闭证据 |
| --- | --- | --- | --- | --- |
{pending_table}

## P-016具体怎样推进

已确定：人员与车辆分型；内部targetId不是展示号；库导入模板、逐行错误和ALL_OR_NOTHING语义已有契约；原始事实、证据和修订谱系必须保留。未确定：哪些来源能证明同一人/车、置信度及阈值、冲突优先级、去重、人工合并拆分的条件和撤销/再计算影响。MAC、IMSI、车牌、人像不能因为同时出现便认定属于同一人。

1. 业务方提供可合法使用的正反例和边界案例；DEV-03编写身份关联规则草案，列出模型来源、证据类型、有效期、冲突和误关联处置。
2. DEV-01编写合并/拆分/修订的事务、权限、版本冲突、谱系、引用任务/数据集的影响方案；不承担发明身份算法的责任。
3. 业务/算法/安全角色确认规则与用途；DEV-04把确认的正反例变成可重复测试。若现有DEV-03不具备算法评测能力，需算法专家或供应方支持，不能把岗位名称当作算法交付能力证明。
4. DEV-03实现候选和证据计算；DEV-01实现授权后的资源写入及历史追溯；DEV-02展示证据、待复核状态和修订流程。当前merge/split等接口有契约、未发现前端运行时引用，不属于已开发页面能力。
5. 保存决策编号、确认人、日期、范围、评测数据/结果；同步主清单、OpenAPI、字典、DDL/迁移、类型、前端和测试。只移除已确认范围的pending，其余子问题保持打开。

无需等所有问题关闭才能开工：可先开发任务底座、授权框架、确定的CRUD、合成样本和已明确的ID集合运算。未决算法、明文输出、真实通知和高风险设备能力继续关闭，不用Mock成功代替正式验收。
''')

module_table='\n'.join(f'| {m[0]} {m[2]} | {m[1]} | {m[3]} | {m[4]} | {m[5]} | {m[6]} |' for m in MODULES)
write('ENGINEERING_ASSIGNMENTS.md',f'''
# 工程师任务与交接边界修订版

日期：2026-10-10。保留原分工M01–M15和四类角色。原始PDF决定业务范围，本文件细化工程实现责任；未知业务规则仍见[PENDING_REGISTER.md](PENDING_REGISTER.md)。配套Word为同目录《多维数据分析平台开发模块与人员分工（修订版）.docx》，修改原文件的副本，原件保持完整。

## 统一分工规则

- 模块主责是交付协调人，负责方案和跨角色验收闭环；不表示独自编写模块全部代码。每个实现工作包只有一个实现责任人。
- DEV-01是**全部公开HTTP接口**、鉴权/数据范围、事务、任务协议、资源写入和迁移发布的实现及契约合并责任人；DEV-03可以提交查询/SQL/算法实现，但不另起一套公开REST认证后端。
- DEV-03负责源协议的数据接入适配、标准化/去重、检索SQL及数据投影、指标、分析/集合计算、身份候选、模型与对应数据Worker。给DEV-01可执行查询/计算模块和样例，不只给一份字段说明。
- DEV-02负责全部前端页面及API适配，沿用Vue 3 + TypeScript + Vite + Element Plus + Pinia + Axios + ECharts + OpenLayers。后端工程师无需重复开发人员库、设备或设置页面。
- DEV-04负责环境、CI/CD、质量门禁、测试数据管理、集成/安全/性能/恢复证据和发布操作；缺陷由对应代码所有者修复，业务算法结论由业务/算法角色确认。
- 公共任务租约/状态机/重试协议由DEV-01实现；数据任务执行函数归DEV-03；通知发送、设备下发等外部副作用执行函数归DEV-01；DEV-04部署所有进程。Worker是运行容器，不是把所有Python代码交给DEV-03的理由。
- 数据库DDL/版本化迁移由DEV-01统一合并发布；DEV-03提交capture/analysis/指标相关变更与查询计划；DEV-04执行空库、升级库及恢复门禁。按受控写入接口授权Worker，禁止双方绕过版本/租约协议直接覆盖资源。

## 逐模块唯一主责和实现拆分

| 模块 | 交付协调主责 | DEV-01后端实现 | DEV-03数据实现 | DEV-02前端实现 | DEV-04验证关口 |
| --- | --- | --- | --- | --- | --- |
{module_table}

## 最容易重叠的责任已这样修正

M04只统筹系统首页，M13统筹运维首页及设备分布；两者复用DEV-03同一套指标，DEV-01包装API。M10负责通知引擎及运维通知后端，DEV-02负责其页面，M13无需另写通知服务。M11继续由DEV-03统筹数据语义和算法，人员库/档案公开API与事务归DEV-01。四类源数据接入协议归DEV-03；对设备下发参数/复位的命令协议归DEV-01；厂商必须提供两类协议。常驻模型仅实现原始PDF要求的RESIDENCE，不把旧分工“轨迹/频次扩展”当成本期要求。

## 必须交接的可执行产物

| 交接号 | 交付方→接收方 | 输入 | 必须交付 | 接收标准 |
| --- | --- | --- | --- | --- |
| H01 | DEV-01→DEV-02/03/04 | 当前OpenAPI和已确认规则 | 固定版本API骨架、错误/权限/状态、启动说明、合成样例 | 类型生成/响应校验通过；有401/403/409/422样例 |
| H02 | DEV-03→DEV-01/04 | 厂商协议、脱敏正反例 | 映射版本、四源解析器、去重键、隔离原因、可重复导入包 | 同源重放不重计；坏行隔离；时间/坐标可核对 |
| H03 | DEV-03→DEV-01 | 标准事实表与授权过滤规范 | 参数化SQL/查询模块、DTO投影、分页和索引计划 | 跨场所/跨类型反例通过；稳定游标；未授权明文不进入投影 |
| H04 | DEV-01→DEV-03 | 统一任务DDL/协议 | 领取/续租/取消/重试/发布接口、租约代次、输入快照格式 | Worker崩溃可恢复；旧租约发布返回失败；不公开半成品 |
| H05 | DEV-03→DEV-01/02/04 | 已签认指标和算法、固定输入 | 可执行计算函数、参数/算法版本、结果和证据、运行样例 | 同输入可复现；空值/取消/失败可解释；效果达到确认标准 |
| H06 | DEV-01→DEV-02 | 可运行API与样例数据 | 同源/api/v1地址、Cookie/CSRF、分页、ETag、文件、配置开关 | 浏览器真实请求入库并读回；服务停止出现真实错误；无Mock回退 |
| H07 | DEV-01→DEV-04 | 通知/设备适配实现与厂商测试端 | 通道/命令超时、UNKNOWN回执对账、逐台结果和密钥注入说明 | 成功项免重发；撤销不谎称撤回外部动作；无凭据入库源码 |
| H08 | DEV-02→DEV-04 | 真实API与已确认页面范围 | 可构建前端、HTTP构建配置、按钮/路由定位和业务流程 | 登录→搜索→地图→任务→文件；分页/过期/越权与多点框选回归 |
| H09 | DEV-01/03→DEV-04 | DDL和旧库快照或空库说明 | 版本化迁移/校验、数据修复与向前恢复步骤 | 空库/升级库核对；引用和文件一致；不能只回退镜像 |
| H10 | DEV-04→全体/项目负责人 | 固定版本、范围、目标环境 | 测试证据、缺陷清单、部署/备份恢复/回滚及移交包 | T01–T27按适用范围通过；业务、安全、运维各自确认 |

## 每位工程师现在就能领的首批任务

DEV-01：创建Spring Boot正式工程；实现统一响应/TraceId与配置、数据库迁移及权限骨架；先提供登录/current-user/菜单/配置/场所和设备的最小真实闭环；实现任务与文件底座后再扩展业务API。认证按当前Cookie拟议契约联调，P-001确认前不宣称认证生产冻结。

DEV-03：先整理四源与状态样本、映射/去重/隔离规范；编写可重复合成数据导入器与参数化检索投影；提交首页/报表指标定义表；按P-011/P-016/P-017分别形成算法提案、样例和评测，不直接用前端Mock算法。

DEV-02：保留现有原型；与DEV-01验证HTTP真实登录/配置/CRUD/检索；补充会话续期接入、后端任务状态刷新策略及已签认但尚无运行时引用的页面动作。66个仅契约接口必须逐项决定调用方和本期范围，健康探针/接入管理API不必强行增加菜单。

DEV-04：建立DEV/TEST隔离环境与CI；固定当前283项单测等前端证据基线；为第一个真实API建立契约、越权、写入/读回、迁移及失败回归；整理容量和恢复输入。当前API样例拦截测试不能计为真实集成通过。

业务/项目负责人不计入四名工程师产能：负责范围与P项签认、厂商协调、数据/模型许可、人员可用性、验收数据和目标；工程师负责技术提案与实现，不能替业务方补造规则。
''')

group_counts={}
for r in api['rows']:
 entry=group_counts.setdefault(r['module'],[0,0,0]);entry[0]+=1;entry[1]+=bool(r['references']);entry[2]+=r['operationPending']
group_table='\n'.join(f'| {k} | {v[0]} | {v[1]} | {v[0]-v[1]} | {v[2]} |' for k,v in group_counts.items())
write('FRONTEND_API_HANDOFF.md',f'''
# 当前前端接口对接清单和后端交接

核对日期：2026-10-10；唯一机器契约[openapi.yaml](../../contracts/api/openapi.yaml)，版本1.0.0-dev.2，SHA256 `{api['summary']['openapiSha256']}`。

当前契约193个操作、581个Schema；65个操作自身标pending。对当前src的TypeScript AST、条件变量和Vue模板检查发现127个操作有运行时静态引用，66个未发现。**静态引用只证明代码出现，不证明按钮可达、业务全部实现或真实后端联调通过。** 逐操作的参数、请求/成功响应Schema、权限、pending、代码文件/行号和交接责任见[JSON](audit/api-inventory.json)；可分发筛选的[CSV](audit/api-inventory.csv)。[66项明细](audit/contract-only-operations.md)保留健康/接入/高级维护等契约，不直接删接口。

| 分配领域 | 全部操作 | 前端静态引用 | 仅契约 | 操作级pending |
| --- | --- | --- | --- | --- |
{group_table}

所有公开API由DEV-01负责实现和统一鉴权，数据查询/计算部分按ENGINEERING_ASSIGNMENTS的H02–H05交接给DEV-03；代码可以同仓协作，不要求跨语言多建公开服务。路径族仅作排期分组；例如批量删除会分派到多个领域，exportMetrics由DEV-03计算指标，认证及通用任务不应被误认为同一业务模块。

## 页面动作怎样调用服务

| 用户动作或页面 | 当前前端主要operationId | 后端必须完成 | 数据工程师必须完成 |
| --- | --- | --- | --- |
| 登录、初始化、退出 | getCsrfToken/login/getMe/getClientConfig/listMenuCatalog/listPermissionCatalog/logout | 会话、CSRF/Origin、用户/范围、配置开关 | 提供合成或批准脱敏的测试数据 |
| 系统首页/运维首页 | getDashboardSummary/getOperationSummary | 授权后的聚合响应和局部失败 | 统一指标定义、水位快照、采集量≠入库量 |
| 四类采集检索 | searchCaptures/getCapture/resolveTargets/createContext | 请求验证、范围/脱敏、游标、Context | 参数化检索、类型payload、索引/投影 |
| 人像上传搜图 | uploadFile/getFileAsset/createFaceSearchTask/getTask/listFaceSearchResults | 文件校验/受权、任务/结果API | 合法模型适配、阈值/上限、候选与评测 |
| 地图/轨迹/热力 | getDeviceMapFeatures/queryTrajectory/getMapHeatmap/getMapAreaSummary | bbox/点数/时间窗和访问范围 | 空间聚合、WGS84转换、按目标排序轨迹 |
| 伴随/碰撞/区域统计 | createAnalysisTask/reviseAnalysisTask/listAnalysisTasks/listAnalysisResults | 任务资源、幂等、范围及发布 | 三类计算和解释证据；保留输入/版本 |
| 数据集/集合计算 | createDataset/importDatasetItems/listDatasetItems/previewSetTask/createSetTask/getSetTask | 数据集版本、引用保护、导入任务 | ID集合运算、导入解析、固定版本 |
| 人员库/图码车码/模型 | listPersonLibraries/createLibraryPerson/createProfileCorrection/searchProfiles/listProfileTargets/listProfileEvidence/listTrendModels/createModelParameterVersion/listTrendModelResults | 库/档案资源、事务/权限/谱系、版本 | 去重/关联候选/证据、RESIDENCE执行 |
| 布控/告警/运维通知 | createWatchRule/activateWatchRule/listAlerts/getWatchStatistics/createNotificationTask/listPlatformMessages | 规则与计划状态、告警去重、可靠投递/回执 | 匹配候选/证据及结束后统计 |
| 场所/设备/参数 | listSites/listDevices/getDeviceParameters/updateDeviceParameters/createDeviceCommand | CRUD/分组、命令下发、逐台回执、审批 | 状态标准化、历史采集/质量数据 |
| 报表/受控导出 | createReport/getReportPreview/downloadReport/createExport/getTask/downloadFile | 任务、文件权限、过期和下载 | 八类公式/预览冻结、Excel/ZIP生成 |
| 用户/角色/日志 | listUsers/listRoles/searchAuditEvents/searchSmsLogs/deleteResourcesBatch/revokeResourcesBatch | 授权、禁用/引用保护、脱敏审计和批量逐项结果 | 查询投影协作，不另起RBAC |

表中列出典型动作，逐项完整清单以JSON/CSV为准。同步201资源创建与异步202执行任务要按每个操作契约区分；通知计划201不代表消息已送达。

## 代码与实际调用链

以前端“围栏检索→查询”为例：`src/components/QueryForm.vue`触发事件，`src/pages/SearchPage.vue`调用Pinia的`src/stores/platform.ts`；`src/api/gateway.ts::loadPage/searchInput`构造FENCE查询并调用`searchCaptures`；`src/api/client.ts::request`按`operations.ts`查到`POST /captures/search`，校验参数，Axios向本站`/api/v1/captures/search`发送JSON。开发时Vite代理转发到后端；生产由反向代理转发。DEV-01验证会话、权限/场所范围并调用DEV-03提交的检索模块，从正式数据库返回受控DTO；前端校验响应，再经`captureView`与Pinia更新表格。HTTP模式无本地Mock成功回退。

长任务示例：前端POST创建分析→API同事务保存意图和async_task→返回202及taskId→数据Worker按租约领取、计算并检查取消→校验租约后原子发布结果→前端getTask查询状态→SUCCESS后读取领域结果或受控下载文件。任务底座与外部副作用由DEV-01负责，数据计算由DEV-03负责；见H04/H05。当前并非所有页面都自动轮询，任务中心提供手工刷新；正式时效要求须按能力明确实现。

## 对接配置和共同约束

详细启动、请求头和交付步骤见[FRONTEND_INTEGRATION.md](../../contracts/api/FRONTEND_INTEGRATION.md)。开发启动前在PowerShell设置`$env:API_PROXY_TARGET='http://127.0.0.1:8080'`；Vite从process.env读取此代理值，不能只把它写进.env.local。`.env.local`中的`VITE_API_MODE=http`和`VITE_API_BASE_URL=/api/v1`决定前端模式/路径。后端保留`/api/v1`前缀，代理没有rewrite。生产须重新以HTTP模式构建并配置同源反向代理；当前演示dist的Mock包不是正式对接包。

- 当前拟议认证为SID Cookie，withCredentials、CSRF和Origin；P-001仍需正式确认，不能同时加入互相冲突的Bearer方案。renewSession有契约、当前前端未发现调用，续期是后续任务。
- 写请求按契约发送X-CSRF-Token、Idempotency-Key；更新资源按要求携带If-Match或批量项version。用户有意重试必须保留同一逻辑请求的键；当前客户端每次未指定键会生成新值，后端不能单凭前端随机键防止所有重复业务。
- 普通成功响应是code="0"、message、data、traceId；错误使用契约HTTP状态和ErrorCode。文件下载返回受控二进制，上传是multipart，不套用JSON文件响应。
- UUID内部ID、targetType、[from,to)时间范围、RFC3339带时区、WGS84经度在前和稳定游标必须保持一致。按钮隐藏不是后端授权。
- 无pending标记不等于可以绕过项目级权限/安全确认；相关能力开关由后端配置。先以真实数据库实现一个小闭环，再逐领域接通，禁止一次性把全部开关打开。

联调通过标准：浏览器实际HTTP请求→服务端日志traceId→真实数据库写入/查询→Worker结果→刷新后读回；同时验证无权、冲突、失败、取消、超时、文件过期、缩权及服务停止。仅OpenAPI样例拦截不满足这项标准。
''')

wbs_table='\n'.join('| '+' | '.join([w[0],w[1],w[2],w[3]]+[f'{a}–{b}' for a,b in w[4]])+' |' for w in WBS)
sumline=' | '.join(f'{a}–{b}' for a,b in totals)
progress='\n'.join(f'| {m[0]} {m[2]} | '+('前端壳和展示交互已实现；生产认证/权限仍须真实联调' if m[0]=='M02' else '已有前端原型与接口/存储设计，真实统计/空间/数据服务未验证' if m[0] in ['M04','M07','M13'] else '已有前端测试和设计验证历史证据；生产CI/部署/容量/恢复未见实现' if m[0]=='M15' else '已有契约/设计或页面外壳；正式后端/Worker未见实现')+f' | {m[6]} |' for m in MODULES)
write('DEVELOPMENT_PLAN.md',f'''
# 实际进度与后续开发计划

截至2026-10-10，本仓库达到“可演示Vue前端原型＋开发级API契约＋数据库设计验证”阶段，未达到完整平台联调或生产验收。未发现正式backend、worker、deploy、CI工程及可连接的真实业务服务证明；这只评价本目录，不能断言其他电脑/仓库没有代码。不以代码文件数、菜单数或API定义数计算整个平台完成百分比。

## 分模块进度与未完成关口

| 模块 | 当前可证实状态 | 后续验证/交付 |
| --- | --- | --- |
{progress}

前端现有28菜单/11类页面、127个接口运行时静态引用和66个仅契约操作。当前283项单测通过；历史28页冒烟、22组扩展流程、16组HTTP样例流程及11项地图框选验证分别是前端证据，不代表193个真实API已实现。数据库163表/1460列/363外键的既有隔离验证证据见design-validation；空库DDL不是已有旧业务库已迁移。20项P编号仍含未决决定。

## 开工前需要的资料

项目负责人指定业务/安全/数据/设备/运营决策角色，按PENDING_REGISTER提供四源和设备状态协议/样本、去重和坐标/时间依据、角色范围矩阵、合法模型/评测集、指标口径、通知和命令测试通道、部署与数据量/SLO/保留/RPO/RTO。先给最小可用样本即可启动基础闭环；人像、身份关联和高风险能力未确认则单独关闭。人员投入、现有后端代码、厂商SDK、历史库及正式合同范围尚需核实，不能假设都已具备。

## 可领任务的工作量基线

以下是工程估算，不是已投入工时或承诺工期。假设四名工程师各一人全职、熟悉当前主栈，覆盖193操作所需公开服务及本期确认的计算/页面/运维能力；有可用协议、许可模型和测试环境；不包含自研高精度人像模型、长期数据标注、采购、资料等待、移动/大屏/HA扩展。DEV-01列包含其负责的外部副作用Worker；DEV-03列是数据计算Worker。单位为每角色人日，工作包之间有共享实现而非193个独立CRUD工程。

| 包 | 工作内容 | 模块 | 依赖/待确认 | DEV-01后端 | DEV-03数据 | DEV-02前端 | DEV-04质量部署 |
| --- | --- | --- | --- | --- | --- | --- | --- |
{wbs_table}
| 合计 | 后续新增工作 | 全范围情景 | 不能四列简单并行相加当日历工期 | {sumline} |

按每人每周4个有效交付人日估算（余下时间用于评审、协调和交接），DEV-01的105–157人日约需27–40周；DEV-03的104–158人日约26–40周，两者构成资源瓶颈。考虑跨角色依赖和末期验收窗口，建议**完整确认范围按34–48周情景做资源计划**，不是确定交付日，也不包含外部等待或新算法研发。若只保留当前前端直接使用的能力、缩减算法/厂商范围，应在逐操作清单确认后重估，不能直接按127/193比例缩时。原计划14周与本次全范围情景不匹配，需要调整范围、人员或日期之一。

## 阶段与可验收里程碑

| 阶段 | 建议窗口（自投入/必要资料到位起） | 实际交付和退出关口 |
| --- | --- | --- |
| G0 规则与项目启动 | 第1–2周；后续未决事项滚动关闭 | 范围/决策人、最小样本、验收矩阵；明确高级接口是否本期；每个P项有责任与期限 |
| G1 首条真实纵向链路 | 第3–10周 | 登录/范围→一个源的检索→地图→受控导出/任务；真实DB读回、401/403/409和失败通过；并行做接入样本和CI |
| G2 数据和管理基础 | 第11–20周 | 四源标准化、稳定检索、管理/版本、首页共享快照、基本报表；回放不重计、权限与版本可测 |
| G3 分析及外部能力增量 | 第21–36周；部分可提前并行 | 集合/研判、人员档案/常驻、布控通知、真实设备命令逐包交付；未签认能力不作为完成 |
| G4 加固与交付 | 第34–48周情景内按实际完成进入 | 真实跨模块T01–T27、升级迁移、安全/容量、备份恢复、发布回滚、操作文档和批准范围验收 |

窗口是资源平衡安排，不是把全部模块等到某一周才开工；G1只选最小纵向切片，不等于W01–W06全部完成。每两周演示可运行增量并更新剩余人日；G1后用实际吞吐校准区间，新增/变更规则同步契约和计划。关键链：授权/任务底座→标准事实→检索/Context→分析/布控→真实通知/命令→生产验收。

## 第一轮两周分配

DEV-01：正式API骨架、错误/TraceId、空库迁移及会话权限最小提案；交H01/H04草案。DEV-03：四源字段映射和可重复合成样本、去重正反例、检索SQL原型；交H02/H03草案。DEV-02：HTTP启动/配置、真实登录和一个列表适配、标明续期与高级动作缺口；DEV-04：DEV/TEST与CI骨架、合同/范围差异登记、真实API与迁移测试框架。第二周末验可运行骨架及样本，不把待确认规则强行实现。

仓库“合同修订工作”的合同副本正文写60个工作日、附件还有完整交付及性能/第三方测试条件；这些不是后端已完成的证据，也不能从副本推断已签署生效。本计划不修改费用、日期或法律条款。项目负责人应先核实正式签署版本与本期范围，再协调工程计划；若必须维持60工作日，需要另行做明确缩围和资源测算，不能用原型截图证明全平台完工。
''')

write('README.md','''
# 项目审查与工程任务交接

本目录是2026-10-10针对现有原型、未决事项、工程师边界和后续开发计划的核对结果。业务要求以用户指定的《平台功能设计（原始）.pdf》为准；本轮不替业务/安全/算法负责人签认未知规则。

| 交付物 | 用途 |
| --- | --- |
| [待确认事项定位与责任](PENDING_REGISTER.md) | P-016在哪、21项主清单和关闭证据 |
| [前端接口交接](FRONTEND_API_HANDOFF.md) | 193操作/127运行时引用/66仅契约，调用链和后端/数据交接 |
| [工程师任务修订版](ENGINEERING_ASSIGNMENTS.md) | 四角色、15模块、10个交接及首批任务 |
| 多维数据分析平台开发模块与人员分工（修订版）.docx | 可编辑Word；在原分工副本中修正文案并补充执行边界 |
| [进度与开发计划](DEVELOPMENT_PLAN.md) | 真实完成度、12工作包、人日、阶段和首轮两周任务 |
| [全目录审查与结果](PROJECT_AUDIT.md) | 审查覆盖、修复、验证和仍缺的输入 |
| [本轮修改清单](MODIFICATIONS.md) | 新增交付、已修改源码/设计文件及保留原件 |
| [API逐项CSV](audit/api-inventory.csv) / [JSON](audit/api-inventory.json) | 每个接口的方法/路径、参数/Schema、代码引用、主责和pending |
| [全文件清单](audit/workspace-files.jsonl) / [汇总](audit/workspace-inventory-summary.json) | 初始快照22,313文件的分类、大小、时间和SHA256；本轮交付另做最终验证 |

现行工程实现边界以本目录修订版为入口；历史审计、旧原型、OpenAPI副本和合同修订资料保留为证据，不用于推断当前生产完成度。原始PDF和原分工DOCX未覆盖修改。
''')

if __name__=='__main__' and '--docx' in __import__('sys').argv:
 doc=Document(SOURCE)
 doc.paragraphs[1].text='2026-10-10 修订版｜模块边界、交接产物与验收清单'
 replacements={
  '建议顺序':'先完成 M01–M03 最小可运行链路，再按依赖并行推进业务；M15质量与交付贯穿全周期。',
  '主责人负责模块方案、核心实现、接口契约、联调与验收闭环；协作人负责其专业范围内的实现和验证。':'模块主责负责交付协调、方案和验收闭环；每个实现工作包按本修订版明确唯一实现人。全部公开HTTP接口归DEV-01，全部前端页面归DEV-02，数据计算/查询模块归DEV-03，质量与部署归DEV-04。',
  '建立用户认证、Token 生命周期、RBAC、菜单权限、数据权限、登录失效和无权访问处理。':'建立认证会话、RBAC、菜单及数据范围、登录失效和无权处理；当前契约拟议SID Cookie，最终方案和生命周期按P-001确认。',
  '实现登录页、Token 注入/续期、退出、超时跳转、菜单树、路由守卫和按钮/数据权限指令。':'实现登录、认证会话接入/续期、退出、过期处理、菜单、路由守卫和按钮权限；依据P-001最终方案，不另行并存冲突的Bearer Token机制。',
  '设计 API Gateway、服务边界、版本策略、统一响应体、错误码、异常处理和 TraceId 传递。':'实现模块化Spring Boot公开API、反向代理边界、版本、统一响应/错误和TraceId；独立网关不是首版必选运行件。',
  '建立统一日志、审计事件、配置中心、定时任务、消息/事件发布和幂等处理机制。':'实现统一日志/审计、环境配置、调度、任务/事件和幂等协议；首版采用数据库任务，不预设独立配置中心或Kafka。',
  '实现系统首页和运维首页的总览可视化，把设备、采集、告警、分布、数据归属和运营状态在一个入口呈现。':'M04负责系统首页12类总览；运维首页四类区域由M13实现，复用同一指标源，避免两次开发。',
  '实现运营首页的设备在线、采集情况、短信猫状态和采集/异常统计四类区域，支持周/月/季度/年度切换。':'运维首页四类区域及周/月/季/年展示归M13；M04只复用统计定义，不重复实现其页面。',
  '由 DEV-03 提供统计口径与聚合接口，由 DEV-01 负责接口权限、缓存和异常降级。':'DEV-03提供指标定义、可执行聚合与投影；DEV-01实现唯一公开聚合API、权限和异常降级；DEV-02展示。',
  '数据趋势提供模型库、模型参数、标签、匹配规则和有效数据筛选；支持驻留、轨迹、频次等趋势模型扩展。':'首版仅按原始PDF提供RESIDENCE常驻模型及参数/标签/匹配/有效数据筛选；算法阈值按P-017签认，不增加未经要求的轨迹/频次模型。',
  '四类检索的查询、重置、分页和导出行为一致；查询条件可从 URL 或任务上下文恢复。':'四类检索交互一致；敏感条件保存于服务端Context，跨页使用contextId；人像普通导出范围按P-007确认。',
  '角色删除受用户引用保护；用户禁用后 Token、任务创建和系统设置权限立即失效或按策略生效。':'角色删除受引用保护；禁用后的会话及在途任务按P-001/P-021策略失效，服务端每次重新授权。',
  '接入 Metrics、Logs、Trace、Alert 和 TraceId 链路；监控 API QPS/延迟/错误率、Kafka、Redis、数据库和业务指标。':'建设Metrics/Logs/Trace/Alert和TraceId；监控实际部署的API、Worker、数据库、文件和业务指标，只有使用Kafka/Redis时才增加对应监控。',
  '把检索结果上的右键动作沉淀为统一的后端任务和上下文服务，保证复制、统计、轨迹、伴随、布控等动作在不同数据类型上口径一致。':'统一检索结果的右键交互；复制由DEV-02在浏览器执行，查询/计算通过DEV-01公开API调用DEV-03模块，跨页和长任务复用Context及任务服务。',
  '实现复制、统计、轨迹、伴随、检索、布控、位置追踪、区域统计、加入数据集和数据导出动作的统一命令模型。':'DEV-02实现复制/重置及动作入口；DEV-01实现统计、轨迹、伴随、检索、布控、位置追踪、区域统计、加入数据集和导出的API/Context编排，按契约区分同步查询与异步任务。',
  '实现交集、并集、差集、混合运算和数据集详情查看，明确同一目标、同一时间和同一来源的去重规则。':'实现交/并/差、表达式运算和数据集详情；CAPTURE按captureId、TARGET按targetType+targetId判等去重。输入成员类型必须兼容，禁止异型集合直接运算，不以时间/来源另造身份规则。',
  '实现 M09 数据集导入与集合运算，固定去重、交并差和跨类型运算规则。':'实现M09数据集导入与集合运算；按P-012固定ID判等和交并差语法，拒绝异型输入；跨源身份关联另按P-016确认。',
  '实现 M11 人员库、图码/车码档案、数据趋势、标签和模型参数，保留原始证据与置信度。':'为M11实现人员导入解析、身份候选/证据、RESIDENCE常驻模型和版本化参数计算；人员库/档案公开API及事务归DEV-01，页面归DEV-02。',
  '按 M05/M06/M08/M09/M10 的接口契约接入数据检索、右键动作、分析任务、数据集和布控。':'按公开API契约接入所有业务页面，包括检索、右键动作、研判、数据集、布控、人员库/档案、设备管理及设置；各模块协调主责不分割前端实现责任。',
  '维护 DEV/TEST/UAT/PROD 环境、测试数据、脱敏、数据库迁移、容器构建、发布、回滚和版本记录。':'维护DEV/TEST/UAT/PROD环境、测试数据和脱敏；执行DEV-01发布的数据库迁移并验证；负责容器构建、发布、回滚和版本记录。',
  '验证备份恢复、服务降级、高可用、通知失败补偿、数据回放、设备离线和采集异常处置。':'验证备份恢复、服务降级、通知失败补偿、数据回放、设备离线和采集异常；高可用扩展只在批准范围明确纳入后另行估算和验收。',
  '主责角色必须完成本模块的核心实现、联调、问题闭环和验收说明；协作角色不得只提交页面截图或临时代码。':'模块主责组织唯一实现责任人完成核心实现、联调、问题闭环和验收；各角色提交其负责的可运行产物及证据，不以页面截图代替服务或计算交付。',
  '文档依据：平台功能设计（原始）.pdf；项目人员要求.docx。':'文档依据：原始PDF为业务基线，现有OpenAPI、数据库设计及前端源码为工程基线；项目人员要求.docx尚未完成等价性验证。',
 }
 modified=[]
 for p in doc.paragraphs:
  for old,new in replacements.items():
   if old in p.text:
    p.text=p.text.replace(old,new);modified.append(p);break
 for t in doc.tables:
  labels=[c.text for c in t.rows[0].cells]
  widths=[1.15,3.65,1.7,3.0,8.9] if labels[:2]==['ID','模块名称'] else [1.6,3.7,6.6,6.5] if labels[:2]==['编号','角色'] else None
  if widths:
   t.autofit=False
   for col,width in zip(t.columns,widths):col.width=Cm(width)
   for row in t.rows:
    for c,width in zip(row.cells,widths):c.width=Cm(width)
  for row in t.rows:
   if row.cells[0].text=='建议顺序': row.cells[1].text=replacements['建议顺序']
   if row.cells[0].text=='DEV-02':row.cells[-1].text='Vue 3、TypeScript、Vite、Element Plus、Pinia、Axios、ECharts、OpenLayers与GeoJSON'
   if row.cells[0].text=='M04':row.cells[-1].text='系统首页12类总览、钻取和导出；运维首页页面统一归M13，指标源共用'
   if row.cells[0].text=='M13':row.cells[-1].text='唯一运维首页四类区块、设备分布和GIS；复用M04指标与M12报表数据'
   if row.cells[0].text=='M11':row.cells[-1].text='人员库、图码/车码档案、证据/谱系和RESIDENCE；DEV-03统筹语义/计算，DEV-01公开API/事务'
   for c in row.cells:
    if re.fullmatch(r'[M\d\s/→–-]+',c.text) and re.search(r'M\d{2}/M\d{2}',c.text):
     c.text=c.text.replace('/',' / ')
 if hasattr(doc,'add_comment'):
  for p in modified[:4]:
   if p.runs:doc.add_comment(p.runs[0],text='2026-10-10责任边界修订：按用户要求细化公开API、数据计算和页面责任；未决认证/算法规则仍需按P清单确认。',author='项目工程审查',initials='RE')
 doc.add_page_break();doc.add_heading('八 工程实现责任与交接补充',level=1)
 doc.add_paragraph('修订日期 2026-10-10。原M01–M15模块主责保持；模块主责是交付协调责任，下面按实现层明确唯一所有者。原始PDF优先，未知业务规则不因本次责任划分而关闭。')
 for text in [
 'DEV-01负责全部公开HTTP接口、鉴权/范围、资源事务、数据库迁移合并发布，以及统一任务租约/状态机。DEV-03可以提交SQL/查询模块，但不另建公开REST后端。',
 'DEV-03负责标准化/去重、接入源协议、查询投影/索引、指标、分析和集合运算、身份候选/证据、常驻模型及数据Worker执行函数。',
 'DEV-02负责全部页面、地图、组件和HTTP适配，沿用当前Vue技术栈；DEV-04负责测试、环境、CI/CD、部署、监控及恢复证据，代码缺陷仍由原作者修复。',
 '通知发送与设备命令下发的副作用Worker由DEV-01实现；DEV-03计算命中候选、标准化遥测和报表数据；DEV-04部署进程。设备数据接入协议归DEV-03，控制命令协议归DEV-01。',
 'DDL/迁移由DEV-01统一管理；DEV-03提交领域结构和查询计划；DEV-04验空库、升级库和恢复。不得绕过任务租约/版本协议直接覆盖业务资源。',
 '业务/安全/算法/设备/运营负责人负责规则、数据用途、许可和验收决定，不能用四名工程师代替决策人；算法能力不足时需专家或供应方支持。']:
  doc.add_paragraph(text)
 doc.add_heading('逐模块实现边界',level=2)
 for m in MODULES:
  doc.add_heading(m[0]+' '+m[2],level=3)
  doc.add_paragraph(f'交付协调主责 {m[1]}。后端 {m[3]}。数据 {m[4]}。前端 {m[5]}。验证 {m[6]}。')
 doc.add_heading('九 跨角色交接与验收',level=1)
 handoffs=[
 ('H01 API契约','DEV-01→DEV-02/03/04。固定OpenAPI、错误/权限/状态、可启动服务及合成样例；类型生成和401/403/409/422校验通过。'),
 ('H02 接入数据','DEV-03→DEV-01/04。源协议和脱敏样本→映射版本、解析器、去重键和隔离记录；回放不重计、坏行可定位。'),
 ('H03 查询模块','DEV-03→DEV-01。参数化SQL、DTO投影、授权范围谓词、稳定游标和索引计划；跨场所/跨类型反例通过。'),
 ('H04 任务协议','DEV-01→DEV-03。领取/续租/取消/重试/原子发布、租约代次和输入快照；崩溃可恢复、旧租约不能发布。'),
 ('H05 数据计算','DEV-03→DEV-01/02/04。可执行函数、固定输入、参数/算法版本、结果和证据；复现、空值、取消及评测通过。'),
 ('H06 真实HTTP','DEV-01→DEV-02。同源/api/v1、Cookie/CSRF、分页/ETag/配置和文件；浏览器写入后刷新读回，后端停止不回退Mock。'),
 ('H07 外部副作用','DEV-01→DEV-04。通知/命令适配、UNKNOWN回执对账、逐台结果、密钥注入；成功项免重发，撤销不谎称撤回已发送动作。'),
 ('H08 前端交付','DEV-02→DEV-04。HTTP构建、定位标识和业务流程；真实登录→检索→地图→任务→下载及多点框选通过。'),
 ('H09 数据库迁移','DEV-01/03→DEV-04。版本化迁移、校验和向前恢复步骤；空库/升级库核对，数据库与文件一起恢复。'),
 ('H10 验收交付','DEV-04→全体/项目负责人。固定版本的T01–T27适用记录、缺陷、部署/回滚/恢复和移交包；业务、安全、运维各自确认。')]
 for title,body in handoffs:doc.add_heading(title,level=2);doc.add_paragraph(body)
 doc.add_heading('十 当前任务与待确认关闭',level=1)
 for role,body in [
 ('DEV-01','创建正式API工程，先做统一响应/TraceId、空库迁移、登录与范围、场所/设备最小CRUD；随后任务/文件/Context；按数据模块接查询和算法。'),
 ('DEV-03','先交四源/状态样本、映射/去重/隔离及合成导入包，检索SQL和指标字典；分别提交P-011/P-016/P-017提案与评测，不复制前端Mock为算法。'),
 ('DEV-02','真实HTTP登录、配置、CRUD、检索、文件与任务；补会话续期和签认后高级动作。当前127操作有运行时静态引用、66项仅契约，不代表全部真实对接。'),
 ('DEV-04','DEV/TEST、CI、真实API/范围反例、迁移及Worker故障测试；保留现有283单测等前端证据，并区分样例拦截与真实集成。')]:
  doc.add_heading(role,level=2);doc.add_paragraph(body)
 doc.add_heading('P-016身份关联与谱系',level=2)
 doc.add_paragraph('主清单为contracts/api/OPENAPI_PENDING.md，数据库投影为database/pending-design.md；data-dictionary.md第1895行和schema.sql的profile_revision注释是影响位置。P-012仅ID集合语法已收敛，其他20项仍含未决决定。')
 doc.add_paragraph('DEV-03提出身份等价/去重、置信度、冲突和模型证据规则；业务/算法/安全角色确认；DEV-01实现合并/拆分/修订事务、权限、版本、谱系和引用保护；DEV-02展示证据及复核，DEV-04验正反例。不得仅因同时出现将MAC、车牌、IMSI、人像认定为同一人。')
 doc.add_paragraph('关闭必须有确认人/日期/适用范围、正反例与评测、受影响OpenAPI/字典/DDL迁移/前端/测试同步；没有证据不移除pending。原始采集事实不被模型或人工结论无痕覆盖。')
 doc.add_heading('十一 进度基线与后续排期',level=1)
 doc.add_paragraph('当前可证实为可演示Vue前端原型、193操作契约和数据库隔离设计验证；本目录未发现正式API/数据Worker/生产部署工程。28菜单不是生产完成率，HTTP样例拦截不是后端联调。全目录初始快照22,313文件均读取并记录SHA256；第三方依赖20,309文件不算自有开发成果。')
 doc.add_paragraph('完整确认范围的后续估算：DEV-01 105–157人日、DEV-03 104–158人日、DEV-02 41–65人日、DEV-04 79–116人日。假设四角色各一人全职、每周4有效交付人日、协议和许可模型可用；不含外部等待、自研模型/数据标注、移动/大屏或生产HA扩展。建议34–48周情景做资源规划，不是承诺日期；需在首条真实链路后按实际效率校准。')
 doc.add_paragraph('第1–2周确认最小范围、决策人和样本；第3–10周首条真实纵向链路；第11–20周四源与管理/指标基础；第21–36周按签认范围增量接分析、档案、布控/通知和设备命令；完成后在34–48周情景内进入真实容量、安全、迁移、恢复和验收。阶段可交叉，不把全部模块延后开工。')
 doc.add_paragraph('原14周草案和合同副本60工作日不能直接当本范围可行承诺；先核实正式范围、人员与资料到位，再协调范围/资源/日期。本修订不改合同费用或期限。详细12工作包、人日和首轮任务见同目录DEVELOPMENT_PLAN.md。')
 doc.add_heading('十二 修订记录与参考入口',level=1)
 doc.add_paragraph('本次在原分工副本中修复M14/M15顺序笔误、M04/M13运维首页重叠、双公开后端职责、人员库事务/算法边界、可选中间件被写成必选、Token与Cookie冲突表述、模型范围扩张和敏感条件URL建议；追加每角色交付物、交接关口和估算。模块编号和原业务验收卡片保留，未决业务不擅自定值。')
 doc.add_paragraph('配套入口：docs/project-governance/README.md；PENDING_REGISTER.md、FRONTEND_API_HANDOFF.md、ENGINEERING_ASSIGNMENTS.md、DEVELOPMENT_PLAN.md、PROJECT_AUDIT.md及audit/api-inventory.csv。原始分工DOCX和原始PDF未覆盖修改。')
 # Retain original document hierarchy and page geometry; apply neutral readable engineering styles.
 for name in ['Normal','List Bullet','List Number','Title','Heading 1','Heading 2','Heading 3']:
  if name not in doc.styles:continue
  style=doc.styles[name];style.font.name='Microsoft YaHei';style.font.color.rgb=RGBColor(0,0,0)
  style._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
  style.font.underline=False
  if name=='Heading 3':style.font.size=Pt(11);style.paragraph_format.space_before=Pt(8);style.paragraph_format.space_after=Pt(4)
  if name.startswith('Heading'):style.paragraph_format.keep_with_next=True
 for p in doc.paragraphs:
  if p.text.strip() in ['开发任务','交付物与验收重点','执行清单','必须协作']:
   p.paragraph_format.keep_with_next=True
  if p.text.startswith('文档依据：'):
   p.paragraph_format.space_before=Pt(2)
   p.paragraph_format.space_after=Pt(0)
   prev=p._p.getprevious()
   if prev is not None and prev.tag==qn('w:p') and not ''.join(prev.itertext()).strip():
    prev.getparent().remove(prev)
   for r in p.runs:r.font.size=Pt(9)
  for r in p.runs:r.font.color.rgb=RGBColor(0,0,0)
  if p.style.name=='Title':
   props=p._p.get_or_add_pPr()
   for border in list(props.findall(qn('w:pBdr'))):props.remove(border)
   styleprops=p.style.element.find(qn('w:pPr'))
   if styleprops is not None:
    for border in list(styleprops.findall(qn('w:pBdr'))):styleprops.remove(border)
 for t in doc.tables:
  for row in t.rows:
   props=row._tr.get_or_add_trPr();nr=OxmlElement('w:cantSplit');props.append(nr)
   for c in row.cells:
    tcpr=c._tc.get_or_add_tcPr();mar=OxmlElement('w:tcMar')
    for side in ['top','left','bottom','right']:
     e=OxmlElement('w:'+side);e.set(qn('w:w'),'85');e.set(qn('w:type'),'dxa');mar.append(e)
    tcpr.append(mar)
    for p in c.paragraphs:
     for r in p.runs:r.font.name='Microsoft YaHei';r.font.color.rgb=RGBColor(0,0,0);r.font.size=Pt(9)
  header=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(header)
  for c in t.rows[0].cells:
   props=c._tc.get_or_add_tcPr()
   for e in list(props.findall(qn('w:shd'))):props.remove(e)
   shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E8EEF3');props.append(shade)
 doc.save(OUT/'多维数据分析平台开发模块与人员分工（修订版）.docx')
 print(json.dumps({'docx':str(OUT/'多维数据分析平台开发模块与人员分工（修订版）.docx'),'modified_original_paragraphs':len(modified),'modules':len(MODULES),'handoffs':len(handoffs),'totals':totals},ensure_ascii=False))
else:
 print(json.dumps({'markdowns_written':5,'pending_items':len(pending_rows),'modules':len(MODULES),'totals':totals},ensure_ascii=False))
